"""Build the CIO Monthly input template: an Excel workbook whose Export tab writes the CSV the
dashboard reads to build the CIO slides without the published PDF (app/src/lib/cioPackage.ts,
docs/cio-template.md).

Two files, both public-safe:
  app/public/templates/CIO_Monthly_Template.xlsx                 blank, for the team
  app/public/templates/CIO_Monthly_Template_Example.xlsx         filled with the latest PUBLIC
                                                                 report, read from the deck's
                                                                 data block (app/public/deck);
                                                                 a stable name, so links hold
Inputs are the yellow cells; Checks repeats the report's own identities (weights to 100%,
categories to the total, 14 bins to 120 months, DM + EM to 100%); Export is formulas only.

Before LACERA publishes a report its figures are confidential: the filled template and its CSV
stay on the team's systems and are read in the browser only. Nothing here runs on real data.

Usage:  python tools/make_cio_template.py
QA (desktop Excel): python tools/qa_cio_template.py
"""

from __future__ import annotations

import datetime as dt
import json
import pathlib
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parents[1]
DECK = ROOT / "app" / "public" / "deck" / "index.html"
OUT = ROOT / "app" / "public" / "templates"
FORMAT = "cio-template-1"

PERIODS = ["1M", "3M", "FYTD", "YTD", "1Y", "3Y", "5Y", "10Y"]
PCOLS = [get_column_letter(3 + i) for i in range(len(PERIODS))]  # C … J
CATS = [("GROWTH", "growth", "Growth"), ("CREDIT", "credit", "Credit"),
        ("RAIH", "ra", "Real Assets & Inflation Hedges"), ("RRM", "rrm", "Risk Reduction & Mitigation")]
BIN_LABELS = ["≤ -6", "-6 to -5", "-5 to -4", "-4 to -3", "-3 to -2", "-2 to -1", "-1 to 0",
              "0 to 1", "1 to 2", "2 to 3", "3 to 4", "4 to 5", "5 to 6", "≥ 6"]
STATS = [("mean", "Mean monthly return"), ("saa", "Policy benchmark (SAA) mean"),
         ("sd", "Standard deviation"), ("min", "Lowest month"), ("max", "Highest month"),
         ("latest", "Latest month")]
STATUS = {"prog": "In progress", "dev": "In development", "info": "For attention",
          "quiet": "Quiet period", "done": "Completed"}
MARKET_SLOTS, MACRO_SLOTS, ITEM_SLOTS, OVERLAY_SLOTS = 20, 8, 30, 4

NAVY = "10233D"
F_TITLE = Font(name="Arial", size=14, bold=True, color=NAVY)
F_HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_BODY = Font(name="Arial", size=10)
F_NOTE = Font(name="Arial", size=9, italic=True, color="55606E")
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_INPUT = PatternFill("solid", fgColor="FFF6CC")
FILL_OK = PatternFill("solid", fgColor="DDF1E4")
FILL_BAD = PatternFill("solid", fgColor="FBDADA")
THIN = Side(style="thin", color="C9D1DA")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# --- layout of a fund tab (rows) ---
R_MV, R_CASH, R_GOD = 4, 5, 6
R_PERF_HEAD = 8
PERF_ROWS = [("TOTAL", "return", "Total Fund — net return"), ("TOTAL", "benchmark", "Total Fund — policy benchmark"),
             ("TOTAL", "hurdle", "Total Fund — actuarial hurdle")]
for code, _, label in CATS:
    PERF_ROWS += [(code, "return", f"{label} — return"), (code, "benchmark", f"{label} — benchmark")]
R_PERF0 = R_PERF_HEAD + 1                                   # 9 … 19
R_ALLOC_HEAD = R_PERF0 + len(PERF_ROWS) + 1                 # 21
R_ALLOC0 = R_ALLOC_HEAD + 1                                 # 22 … 25 categories, 26 OTHER
R_OTHER = R_ALLOC0 + len(CATS)
R_ALLOC_SUM = R_OTHER + 1
R_OVL_HEAD = R_ALLOC_SUM + 2
R_OVL0 = R_OVL_HEAD + 1
R_HIST_HEAD = R_OVL0 + OVERLAY_SLOTS + 1
R_BIN0 = R_HIST_HEAD + 1
R_STAT0 = R_BIN0 + len(BIN_LABELS) + 1
R_GEO_HEAD = R_STAT0 + len(STATS) + 1
R_DM, R_EM, R_MKTS = R_GEO_HEAD + 1, R_GEO_HEAD + 2, R_GEO_HEAD + 3
R_CDM0 = R_MKTS + 2
R_CEM0 = R_CDM0 + 6


def deck_data() -> dict:
    """The latest public report, from the deck's generated data block (JSON literals)."""
    text = DECK.read_text(encoding="utf-8")
    out = {}
    for name in ("PERIODS", "ENT", "MKT", "MACRO", "OPS", "VINTAGE"):
        m = re.search(rf"^  let {name} = (.*);\r?$", text, re.M)
        if not m:
            raise SystemExit(f"{name} not found in the deck's data block")
        out[name] = json.loads(m.group(1))
    if out["PERIODS"] != ["1 M", "3 M", "FYTD", "YTD", "1 Y", "3 Y", "5 Y", "10 Y"]:
        raise SystemExit("the deck's periods changed; update PERIODS here")
    return out


def iso(label: str) -> dt.date:
    return dt.datetime.strptime(label, "%B %d, %Y").date()


def head(ws, row: int, labels: list[str], col: int = 1) -> None:
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font, c.fill, c.alignment = F_HEAD, FILL_HEAD, Alignment(horizontal="center" if i else "left")


def inp(ws, ref: str, value=None, fmt: str | None = None) -> None:
    c = ws[ref]
    if value is not None:
        c.value = value
    c.fill, c.border, c.font = FILL_INPUT, BOX, F_BODY
    if fmt:
        c.number_format = fmt


def title(ws, text: str, note: str | None = None) -> None:
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if note:
        ws["A2"] = note
        ws["A2"].font = F_NOTE


def fund_tab(wb, name: str, e: dict | None) -> None:
    ws = wb.create_sheet(name)
    title(ws, name, "Yellow cells are inputs. Percent as printed (0.1 means 0.1%); money in $ millions.")
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 30
    for col in PCOLS:
        ws.column_dimensions[col].width = 10
    ws.freeze_panes = "C4"

    for r, label in ((R_MV, "Total market value ($ millions)"), (R_CASH, "Cash and equivalents ($ millions)"),
                     (R_GOD, "Growth of $1, trailing 5 years (optional)")):
        ws.cell(row=r, column=1, value=label).font = F_BODY
    inp(ws, f"C{R_MV}", e and e["mv"], "#,##0")
    inp(ws, f"C{R_CASH}", e and e["cash"], "#,##0")
    inp(ws, f"C{R_GOD}", e and e.get("god"), "0.00")

    head(ws, R_PERF_HEAD, ["Performance (% net of fees)", ""] + PERIODS)
    for i, (code, measure, label) in enumerate(PERF_ROWS):
        r = R_PERF0 + i
        ws.cell(row=r, column=1, value=label).font = F_BODY
        if e:
            if code == "TOTAL":
                src = e["total"]["r" if measure == "return" else "b" if measure == "benchmark" else "h"]
            else:
                comp = next(c for c in e["comps"] if c["k"] == dict((a, b) for a, b, _ in CATS)[code])
                src = comp["r" if measure == "return" else "b"]
        for j, col in enumerate(PCOLS):
            inp(ws, f"{col}{r}", e and src[j], "0.0")

    head(ws, R_ALLOC_HEAD, ["Allocation", "Label (other line only)", "Market value $M", "Weight %", "Target %", "Month flow $M"])
    for i, (code, k, label) in enumerate(CATS):
        r = R_ALLOC0 + i
        ws.cell(row=r, column=1, value=label).font = F_BODY
        comp = e and next(c for c in e["comps"] if c["k"] == k)
        inp(ws, f"C{r}", comp and comp["mv"], "#,##0")
        inp(ws, f"D{r}", comp and comp["pct"], "0.0")
        inp(ws, f"E{r}", comp and comp["tgt"], "0.0")
        inp(ws, f"F{r}", comp and comp["flow"], "#,##0")
    ws.cell(row=R_OTHER, column=1, value="Other line (overlays, other assets, unitemized; optional)").font = F_BODY
    other = e and e.get("other")
    inp(ws, f"B{R_OTHER}", other and other["n"])
    inp(ws, f"C{R_OTHER}", other and other["mv"], "#,##0")
    inp(ws, f"D{R_OTHER}", other and other["pct"], "0.0")
    inp(ws, f"F{R_OTHER}", other and other["flow"], "#,##0")
    ws.cell(row=R_ALLOC_SUM, column=1, value="Sum (should match the total and 100%)").font = F_NOTE
    for col, fmt in (("C", "#,##0"), ("D", "0.0"), ("F", "#,##0")):
        c = ws[f"{col}{R_ALLOC_SUM}"]
        c.value = f"=SUM({col}{R_ALLOC0}:{col}{R_OTHER})"
        c.number_format, c.font = fmt, F_NOTE

    head(ws, R_OVL_HEAD, ["Overlay programs (Total Fund only, $ millions)", "Program", "Month gain", "Since inception"])
    ovl = (e and e.get("overlays")) or []
    for i in range(OVERLAY_SLOTS):
        r = R_OVL0 + i
        o = ovl[i] if i < len(ovl) else None
        inp(ws, f"B{r}", o and o["n"])
        inp(ws, f"C{r}", o and o["may"], "#,##0.0")
        inp(ws, f"D{r}", o and o["si"], "#,##0.0")

    head(ws, R_HIST_HEAD, ["Return distribution, last 120 months (% per month)", "Bin", "Months"])
    for i, label in enumerate(BIN_LABELS):
        r = R_BIN0 + i
        ws.cell(row=r, column=2, value=label).font = F_BODY
        inp(ws, f"C{r}", e and e["hist"]["c"][i], "0")
    for i, (k, label) in enumerate(STATS):
        r = R_STAT0 + i
        ws.cell(row=r, column=2, value=label).font = F_BODY
        inp(ws, f"C{r}", e and e["hist"][k], "0.00")

    head(ws, R_GEO_HEAD, ["Geographic exposure (optional)", "", "Share %", "Markets"])
    g = e and e.get("geo")
    has_geo = bool(g and g.get("top"))
    for r, label, share, n in ((R_DM, "Developed markets", "dm", "dmN"), (R_EM, "Emerging / frontier markets", "em", "emN")):
        ws.cell(row=r, column=1, value=label).font = F_BODY
        inp(ws, f"C{r}", g[share] if has_geo else None, "0.0")
        inp(ws, f"D{r}", g[n] if has_geo else None, "0")
    ws.cell(row=R_MKTS, column=1, value="Total markets").font = F_BODY
    inp(ws, f"D{R_MKTS}", g["total"] if has_geo else None, "0")
    for r0, grp, label in ((R_CDM0, "dm", "Top developed-market countries"), (R_CEM0, "em", "Top emerging-market countries")):
        ws.cell(row=r0 - 1, column=1, value=label).font = F_NOTE
        top = [c for c in (g["top"] if has_geo else []) if c[2] == grp]
        for i in range(5):
            c = top[i] if i < len(top) else None
            inp(ws, f"B{r0 + i}", c and c[0])
            inp(ws, f"C{r0 + i}", c and c[1], "0.0")


def build(example: dict | None) -> Workbook:
    wb = Workbook()
    start = wb.active
    start.title = "Start"
    title(start, "CIO Monthly — input template")
    lines = [
        "Fills the CIO Monthly slides on the dashboard without the published PDF.",
        "",
        "1. Fill the yellow cells on Report, Pension, OPEB, Markets, Macro and Items.",
        "   Percent as printed (0.1 means 0.1%); money in $ millions. Leave a cell empty when the report has no figure.",
        "2. Open Checks: every line should say OK.",
        '3. Open Export, then File > Save As > "CSV UTF-8 (Comma delimited)". Excel saves only the Export tab.',
        '4. On the dashboard\'s CIO Monthly tab choose "Open a template file" and pick the CSV.',
        "   It is read in your browser only: nothing is uploaded, and it is cleared when you close or reload the page.",
        "",
        "Before LACERA publishes a report, its figures are confidential. Keep this workbook and its CSV on your",
        "team's systems. Do not use the GitHub \"CIO report\" form for them, and do not send them to AI tools.",
        "",
        f"Format {FORMAT}. Column meanings: docs/cio-template.md in the project on GitHub.",
    ]
    for i, text in enumerate(lines):
        start.cell(row=3 + i, column=1, value=text).font = F_BODY
    start.column_dimensions["A"].width = 110

    v = example and example["VINTAGE"]
    rep = wb.create_sheet("Report")
    title(rep, "Report", "Dates as YYYY-MM-DD. Data through is the month end the fund figures are as of.")
    rep.column_dimensions["A"].width = 44
    rep.column_dimensions["C"].width = 14
    for r, label, val in ((4, "Report (meeting) date", v and iso(v["reportLabel"])),
                          (5, "Data through (month end)", v and iso(v["throughLabel"])),
                          (6, "Market table as of", v and iso(v["marketLabel"]))):
        rep.cell(row=r, column=1, value=label).font = F_BODY
        inp(rep, f"C{r}", val, "yyyy-mm-dd")

    ent = example and example["ENT"]
    fund_tab(wb, "Pension", ent and ent["pension"])
    fund_tab(wb, "OPEB", ent and ent["opeb"])

    mk = wb.create_sheet("Markets")
    title(mk, "Market table (index total returns, %)", "One row per index, grouped as the report groups them.")
    head(mk, 4, ["Group", "Index", "Description"] + PERIODS)
    for col, w in (("A", 28), ("B", 26), ("C", 48)):
        mk.column_dimensions[col].width = w
    rows = [(g["g"], r) for g in (example["MKT"] if example else []) for r in g["rows"]]
    for i in range(MARKET_SLOTS):
        r = 5 + i
        g, row = rows[i] if i < len(rows) else (None, None)
        inp(mk, f"A{r}", g)
        inp(mk, f"B{r}", row and row["n"])
        inp(mk, f"C{r}", row and row["i"])
        for j in range(len(PERIODS)):
            inp(mk, f"{get_column_letter(4 + j)}{r}", row and row["v"][j], "0.0")

    mc = wb.create_sheet("Macro")
    title(mc, "Macro strip", "Label, value as shown on the slide, and a detail line with its source or page.")
    head(mc, 4, ["Label", "Value", "Detail"])
    for col, w in (("A", 40), ("B", 22), ("C", 100)):
        mc.column_dimensions[col].width = w
    macro = example["MACRO"] if example else []
    for i in range(MACRO_SLOTS):
        r = 5 + i
        m = macro[i] if i < len(macro) else None
        inp(mc, f"A{r}", m and m["l"])
        inp(mc, f"B{r}", m and m["v"])
        inp(mc, f"C{r}", m and m["s"])

    it = wb.create_sheet("Items")
    title(it, "Items for attention", "Area, the item as it should read, its status and the report page (optional).")
    head(it, 4, ["Area", "Item", "Status", "Page"])
    for col, w in (("A", 24), ("B", 100), ("C", 16), ("D", 8)):
        it.column_dimensions[col].width = w
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS.values()) + '"', allow_blank=True)
    it.add_data_validation(dv)
    ops = example["OPS"] if example else []
    for i in range(ITEM_SLOTS):
        r = 5 + i
        o = ops[i] if i < len(ops) else None
        inp(it, f"A{r}", o and o["e"])
        inp(it, f"B{r}", o and o["item"])
        inp(it, f"C{r}", o and STATUS[o["st"]])
        inp(it, f"D{r}", o and o["p"], "0")
        dv.add(f"C{r}")

    checks(wb)
    export(wb)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title in ("Export",)
    stamp = dt.datetime.combine(iso(v["reportLabel"]), dt.time()) if v else dt.datetime(2026, 1, 1)
    wb.properties.creator = "tools/make_cio_template.py"
    wb.properties.created = wb.properties.modified = stamp
    return wb


def checks(wb) -> None:
    ws = wb.create_sheet("Checks")
    title(ws, "Checks", "The report's own identities. Every line should say OK before you export.")
    head(ws, 4, ["Check", "Value", "Result"])
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10
    rows = [
        ("Data through is a month end", "=Report!C5", '=IF(AND(ISNUMBER(Report!C5),EOMONTH(Report!C5,0)=Report!C5),"OK","CHECK")'),
        ("Report date is on or after data through", "=Report!C4", '=IF(AND(ISNUMBER(Report!C4),Report!C4>=Report!C5),"OK","CHECK")'),
    ]
    for tab, name in (("Pension", "Pension Fund"), ("OPEB", "OPEB Trust")):
        s = f"'{tab}'!" if " " in tab else f"{tab}!"
        cats = f"{s}C{R_ALLOC0}:C{R_OTHER}"
        core = ",".join(f"{s}{PCOLS[j]}{R_PERF0 + k}" for k in (0, 1) for j in (0, 2, 4))
        pairs = "+".join(
            f"SUMPRODUCT(--(({s}C{R_PERF0 + 3 + 2 * i}:J{R_PERF0 + 3 + 2 * i}=\"\")<>({s}C{R_PERF0 + 4 + 2 * i}:J{R_PERF0 + 4 + 2 * i}=\"\")))"
            for i in range(len(CATS)))
        pairs += f"+SUMPRODUCT(--(({s}C{R_PERF0}:J{R_PERF0}=\"\")<>({s}C{R_PERF0 + 1}:J{R_PERF0 + 1}=\"\")))"
        rows += [
            (f"{name}: total market value and cash entered", f"=COUNT({s}C{R_MV},{s}C{R_CASH})",
             f'=IF(COUNT({s}C{R_MV},{s}C{R_CASH})=2,"OK","CHECK")'),
            (f"{name}: Total Fund return and benchmark for 1M, FYTD and 1Y", f"=COUNT({core})",
             f'=IF(COUNT({core})=6,"OK","CHECK")'),
            (f"{name}: every return has a benchmark for the same period", f"={pairs}",
             f'=IF({pairs}=0,"OK","CHECK")'),
            (f"{name}: category weights add to 100% (±0.3)", f"=SUM({s}D{R_ALLOC0}:D{R_OTHER})",
             f'=IF(ABS(SUM({s}D{R_ALLOC0}:D{R_OTHER})-100)<=0.3,"OK","CHECK")'),
            (f"{name}: category market values add to the total", f"=SUM({cats})",
             f'=IF(AND(ISNUMBER({s}C{R_MV}),ABS(SUM({cats})-{s}C{R_MV})<=MAX(2,0.003*{s}C{R_MV})),"OK","CHECK")'),
            (f"{name}: 14 bin counts add to 120 months", f"=SUM({s}C{R_BIN0}:C{R_BIN0 + 13})",
             f'=IF(AND(COUNT({s}C{R_BIN0}:C{R_BIN0 + 13})=14,SUM({s}C{R_BIN0}:C{R_BIN0 + 13})=120),"OK","CHECK")'),
            (f"{name}: six distribution statistics entered", f"=COUNT({s}C{R_STAT0}:C{R_STAT0 + 5})",
             f'=IF(COUNT({s}C{R_STAT0}:C{R_STAT0 + 5})=6,"OK","CHECK")'),
            (f"{name}: DM + EM shares add to 100% (if geography entered)", f"=SUM({s}C{R_DM}:C{R_EM})",
             f'=IF(COUNT({s}C{R_DM}:C{R_EM})=0,"OK",IF(AND(COUNT({s}C{R_DM}:C{R_EM})=2,ABS(SUM({s}C{R_DM}:C{R_EM})-100)<=0.2),"OK","CHECK"))'),
        ]
    rows.append(("Market table has its as-of date (if indexes entered)", "=COUNTA(Markets!B5:B24)",
                 '=IF(OR(COUNTA(Markets!B5:B24)=0,ISNUMBER(Report!C6)),"OK","CHECK")'))
    rows.append(("Every market index has a group", '=SUMPRODUCT(--(Markets!B5:B24<>""),--(Markets!A5:A24=""))',
                 '=IF(SUMPRODUCT(--(Markets!B5:B24<>""),--(Markets!A5:A24=""))=0,"OK","CHECK")'))
    for i, (label, value, result) in enumerate(rows):
        r = 5 + i
        ws.cell(row=r, column=1, value=label).font = F_BODY
        ws.cell(row=r, column=2, value=value).font = F_BODY
        c = ws.cell(row=r, column=3, value=result)
        c.font, c.alignment = Font(name="Arial", size=10, bold=True), Alignment(horizontal="center")
    last = 4 + len(rows)
    ws.conditional_formatting.add(f"C5:C{last}", CellIsRule(operator="equal", formula=['"OK"'], fill=FILL_OK))
    ws.conditional_formatting.add(f"C5:C{last}", CellIsRule(operator="equal", formula=['"CHECK"'], fill=FILL_BAD))
    ws.cell(row=last + 2, column=1, value="All checks").font = Font(name="Arial", size=10, bold=True)
    c = ws.cell(row=last + 2, column=3, value=f'=IF(COUNTIF(C5:C{last},"OK")={len(rows)},"OK","CHECK")')
    c.font, c.alignment = Font(name="Arial", size=10, bold=True), Alignment(horizontal="center")
    ws.conditional_formatting.add(f"C{last + 2}", CellIsRule(operator="equal", formula=['"OK"'], fill=FILL_OK))
    ws.conditional_formatting.add(f"C{last + 2}", CellIsRule(operator="equal", formula=['"CHECK"'], fill=FILL_BAD))


def iso_text(ref: str) -> str:
    """ISO date text from a date cell, independent of the Excel locale's format codes."""
    return f'IF({ref}="","",YEAR({ref})&"-"&TEXT(MONTH({ref}),"00")&"-"&TEXT(DAY({ref}),"00"))'


def export(wb) -> None:
    ws = wb.create_sheet("Export")
    cols = ["section", "entity", "item", "measure", "period", "value", "text"]
    for i, name in enumerate(cols):
        ws.cell(row=1, column=1 + i, value=name)
    rows: list[list] = []
    val = lambda ref: f'=IF({ref}="","",{ref})'  # noqa: E731 - a blank input stays blank
    rows.append(["report", "", "format", "", "", "", FORMAT])
    for item, ref in (("report_date", "Report!C4"), ("data_through", "Report!C5"), ("market_as_of", "Report!C6")):
        rows.append(["report", "", item, "", "", "", "=" + iso_text(ref)])
    for tab, ent in (("Pension", "pension"), ("OPEB", "opeb")):
        s = f"{tab}!"
        for r, measure in ((R_MV, "market_value"), (R_CASH, "cash"), (R_GOD, "growth_of_dollar")):
            rows.append(["fund", ent, "TOTAL", measure, "", val(f"{s}C{r}"), ""])
        for i, (code, measure, _) in enumerate(PERF_ROWS):
            for j, period in enumerate(PERIODS):
                rows.append(["performance", ent, code, measure, period, val(f"{s}{PCOLS[j]}{R_PERF0 + i}"), ""])
        for i, (code, _, _) in enumerate(CATS):
            r = R_ALLOC0 + i
            for col, measure in (("C", "market_value"), ("D", "weight"), ("E", "target"), ("F", "flow")):
                rows.append(["allocation", ent, code, measure, "", val(f"{s}{col}{r}"), ""])
        rows.append(["allocation", ent, "OTHER", "label", "", "", val(f"{s}B{R_OTHER}")])
        for col, measure in (("C", "market_value"), ("D", "weight"), ("F", "flow")):
            rows.append(["allocation", ent, "OTHER", measure, "", val(f"{s}{col}{R_OTHER}"), ""])
        for i in range(OVERLAY_SLOTS):
            r = R_OVL0 + i
            for col, measure in (("C", "month_gain"), ("D", "since_inception")):
                rows.append(["overlay", ent, val(f"{s}B{r}"), measure, "",
                             f'=IF(OR({s}B{r}="",{s}{col}{r}=""),"",{s}{col}{r})', ""])
        for i in range(len(BIN_LABELS)):
            rows.append(["histogram", ent, f"BIN_{i:02d}", "count", "", val(f"{s}C{R_BIN0 + i}"), ""])
        for i, (k, _) in enumerate(STATS):
            rows.append(["histogram", ent, "STAT", k, "", val(f"{s}C{R_STAT0 + i}"), ""])
        for r, item in ((R_DM, "DM"), (R_EM, "EM")):
            rows.append(["geography", ent, item, "share", "", val(f"{s}C{r}"), ""])
            rows.append(["geography", ent, item, "markets", "", val(f"{s}D{r}"), ""])
        rows.append(["geography", ent, "TOTAL", "markets", "", val(f"{s}D{R_MKTS}"), ""])
        for r0, grp in ((R_CDM0, "dm"), (R_CEM0, "em")):
            for i in range(5):
                r = r0 + i
                rows.append(["country", ent, val(f"{s}B{r}"), "share", "",
                             f'=IF(OR({s}B{r}="",{s}C{r}=""),"",{s}C{r})', f'=IF({s}B{r}="","","{grp}")'])
    for i in range(MARKET_SLOTS):
        r = 5 + i
        name = f"Markets!B{r}"
        rows.append(["market", "", val(name), "group", "", "", f'=IF({name}="","",Markets!A{r})'])
        rows.append(["market", "", val(name), "description", "", "", f'=IF({name}="","",Markets!C{r})'])
        for j, period in enumerate(PERIODS):
            cell = f"Markets!{get_column_letter(4 + j)}{r}"
            rows.append(["market", "", val(name), "return", period, f'=IF(OR({name}="",{cell}=""),"",{cell})', ""])
    for i in range(MACRO_SLOTS):
        r = 5 + i
        label = f"Macro!A{r}"
        rows.append(["macro", "", val(label), "value", "", "", f'=IF({label}="","",Macro!B{r})'])
        rows.append(["macro", "", val(label), "detail", "", "", f'=IF({label}="","",Macro!C{r})'])
    for i in range(ITEM_SLOTS):
        r = 5 + i
        item = f"Items!B{r}"
        rows.append(["attention", "", f'=IF({item}="","",Items!A{r})', f'=IF({item}="","",Items!C{r})', "",
                     f'=IF(OR({item}="",Items!D{r}=""),"",Items!D{r})', f'=IF({item}="","",{item})'])
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            if value != "":
                ws.cell(row=2 + i, column=1 + j, value=value)
    for i, w in enumerate((12, 9, 30, 16, 7, 12, 40)):
        ws.column_dimensions[get_column_letter(1 + i)].width = w
    ws.freeze_panes = "A2"


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    data = deck_data()
    blank = OUT / "CIO_Monthly_Template.xlsx"
    build(None).save(blank)
    example = OUT / "CIO_Monthly_Template_Example.xlsx"
    build(data).save(example)
    for p in (blank, example):
        print(f"wrote {p.relative_to(ROOT)} ({p.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
